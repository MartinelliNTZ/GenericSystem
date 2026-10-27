# -*- coding: utf-8 -*-
"""
ResumoOsService — Planilha-resumo da base PRIVADA (VERRA_DATA)
=============================================================
Mantem o XLSX de acompanhamento da base privada, gravado na RAIZ da
``VERRA_DATA`` (a mesma pasta dos completões ``Consolidado`` e das pastas de OS):

    <BASE_DADOS>/VERRA_DATA/resumo_os.xlsx

Aba ``Resumo``, com o cabeçalho e UMA LINHA POR OS:

    OS | ÁREA FINAL | ZM

- **OS**: chave OFICIAL da OS como TEXTO (``050``, ``169A``) — o zero à esquerda
  é preservado (a célula sai com formato de texto).
- **ÁREA FINAL**: soma de ``AREA_HA`` de TODAS as feições do shape de LIMITES da
  OS (em hectares).
- **ZM**: data em que as zonas de manejo da OS foram inseridas no sistema
  (``inserted_at`` do registro ``zonas_manejo`` da OS no banco oficial), gravada
  como DATA real (``DD/MM/YYYY``).

A planilha é REGERADA por inteiro a cada IMPORTAR/ATUALIZAR/ATUALIZAR BASE
(nunca no VALIDAR) e a gravação é ATÔMICA (temporário + substituição). Quando o
valor da ZM não vem do banco, a data já gravada na planilha é PRESERVADA — o
acompanhamento não se perde.

Uso:
    from plugins.limites_db.ResumoOsService import ResumoOsService

    ResumoOsService.write(root, [(os_key, area_ha, inserted_at), ...], tool_key)
"""

from __future__ import annotations

import os
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.enum.ToolKey import ToolKey
from utils.BaseUtil import BaseUtil

_TOOL_KEY = ToolKey.LIMITES_DB.value


class ResumoOsService(BaseUtil):
    """Planilha-resumo da base privada: OS x area final x data da ZM."""

    FILE_NAME = "resumo_os.xlsx"
    SHEET_NAME = "Resumo"
    HEADERS: Tuple[str, ...] = ("OS", "ÁREA FINAL", "ZM")

    # Formatos das células (Excel): OS como TEXTO, área numérica e ZM como DATA.
    TEXT_FORMAT = "@"
    AREA_FORMAT = "#,##0.00"
    DATE_FORMAT = "DD/MM/YYYY"
    COLUMN_WIDTHS: Tuple[int, ...] = (14, 16, 14)

    AREA_DECIMALS = 2
    # Formatos aceitos na leitura da data da ZM (banco e planilha anterior).
    DATE_INPUTS: Tuple[str, ...] = (
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%d/%m/%Y",
    )

    # ── API ─────────────────────────────────────────────────────────

    @classmethod
    def path(cls, root: str | Path) -> Path:
        """Caminho da planilha na RAIZ da base privada (``VERRA_DATA``)."""
        return Path(root) / cls.FILE_NAME

    @classmethod
    def write(
        cls,
        root: str | Path,
        rows: Iterable[Sequence[Any]],
        tool_key: str = _TOOL_KEY,
    ) -> Optional[str]:
        """
        Regera a planilha-resumo e devolve o caminho gravado (``None`` sem dados).

        Args:
            root: Raiz da base privada (``VERRA_DATA``).
            rows: Sequência de ``(OS, área final em ha, data da ZM)`` — a área
                pode ser ``None`` e a data pode vir vazia (nesse caso a data já
                gravada na planilha é preservada).
            tool_key: Chave da ferramenta para logging.

        Returns:
            Caminho do XLSX gravado (``None`` quando não há nenhuma OS).
        """
        logger = cls._get_logger(tool_key, cls.__name__)
        target = cls.path(root)
        previous = cls._read_dates(target, logger)
        data = [cls._row(row, previous) for row in rows]

        if not data:
            logger.info(
                "Planilha-resumo nao gerada (base vazia)",
                code="RESUMO_EMPTY",
                root=str(root),
            )
            return None

        cls._save(cls._build_sheet(data), target, logger)
        logger.info(
            "Planilha-resumo atualizada",
            code="RESUMO_OK",
            path=str(target),
            n_os=len(data),
        )
        return str(target)

    # ── Linhas ──────────────────────────────────────────────────────

    @classmethod
    def _row(cls, row: Sequence[Any], previous: Dict[str, Any]) -> Tuple[Any, ...]:
        """Linha pronta: OS como texto, área arredondada e data da ZM resolvida."""
        values = list(row) + [None, None, None]
        os_key = str(values[0] or "").strip()
        area = cls._round_area(values[1])
        return os_key, area, cls._date(values[2], previous.get(os_key))

    @classmethod
    def _round_area(cls, area: Any) -> Optional[float]:
        """Área em hectares com as casas decimais da planilha (``None`` sem área)."""
        if area is None:
            return None
        try:
            return round(float(area), cls.AREA_DECIMALS)
        except (TypeError, ValueError):
            return None

    @classmethod
    def _date(cls, value: Any, fallback: Any) -> Any:
        """Data da ZM: valor atual (banco); sem ele, a data já gravada na planilha."""
        return cls._parse_date(value) or cls._parse_date(fallback) or ""

    @classmethod
    def _parse_date(cls, value: Any) -> Optional[datetime]:
        """Converte ISO/``DD/MM/AAAA``/``date`` em ``datetime`` (``None`` sem valor)."""
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime(value.year, value.month, value.day)

        text = str(value or "").strip()
        if not text:
            return None

        for fmt in cls.DATE_INPUTS:
            parsed = cls._try_format(text, fmt)
            if parsed is not None:
                return parsed

        return None

    @staticmethod
    def _try_format(text: str, fmt: str) -> Optional[datetime]:
        """Tenta ler ``text`` no formato ``fmt`` (``None`` quando não casa)."""
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            return None

    # ── Planilha ────────────────────────────────────────────────────

    @classmethod
    def _build_sheet(cls, data: Sequence[Sequence[Any]]) -> Workbook:
        """Monta a aba do resumo com o cabeçalho em negrito, formatos e larguras."""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = cls.SHEET_NAME
        sheet.append(list(cls.HEADERS))

        for cell in sheet[1]:
            cell.font = Font(bold=True)

        for row in data:
            sheet.append(list(row))

        cls._format_columns(sheet)
        return workbook

    @classmethod
    def _format_columns(cls, sheet: Any) -> None:
        """Aplica o formato de cada coluna (texto/número/data) e a largura."""
        for index, width in enumerate(cls.COLUMN_WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width

        formats = (cls.TEXT_FORMAT, cls.AREA_FORMAT, cls.DATE_FORMAT)
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for index, number_format in enumerate(formats):
                row[index].number_format = number_format

    @classmethod
    def _read_dates(cls, target: Path, logger: Any) -> Dict[str, Any]:
        """Datas da ZM já gravadas na planilha atual (``{}`` sem arquivo/ilegível).

        O arquivo é SEMPRE fechado: no modo ``read_only`` o openpyxl mantém o
        handle aberto e a substituição atômica seguinte falharia (Windows).
        """
        if not target.is_file():
            return {}

        workbook = None
        try:
            workbook = load_workbook(target, read_only=True, data_only=True)
            sheet = workbook[cls.SHEET_NAME]
            return {
                str(row[0]).strip(): row[2]
                for row in sheet.iter_rows(min_row=2, values_only=True)
                if row and len(row) > 2 and row[0] and row[2]
            }
        except Exception as e:
            logger.warning(
                "Planilha-resumo ilegivel — datas gravadas nao puderam ser preservadas",
                code="RESUMO_READ_ERR",
                path=str(target),
                error=str(e),
            )
            return {}
        finally:
            if workbook is not None:
                workbook.close()

    @staticmethod
    def _save(workbook: Workbook, target: Path, logger: Any) -> None:
        """Grava a planilha de forma ATÔMICA (temporário na mesma pasta + troca)."""
        tmp = target.with_name(f"{target.name}.tmp")
        try:
            workbook.save(str(tmp))
            os.replace(str(tmp), str(target))
        finally:
            if tmp.exists():
                tmp.unlink()

        logger.debug("Planilha-resumo gravada", code="RESUMO_SAVE", path=str(target))

